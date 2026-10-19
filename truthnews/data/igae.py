"""IGAE: extracto mensual de ejecucion del Presupuesto del Estado (Administracion General del Estado).

La IGAE publica cada mes un Excel con la ejecucion acumulada (hoja GTOS 001: gastos por seccion o
ministerio; GTOS 002: inversiones reales, capitulo 6; ING 002: ingresos). Se lee la pagina publica del
extracto, se toma el Excel mas reciente y se citan su URL y el periodo. Importes en miles de euros, como
los publica la IGAE; los porcentajes se calculan aqui a partir de los importes.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from urllib.parse import urljoin

import httpx

from truthnews.data import client

BASE = "https://www.igae.pap.hacienda.gob.es"
PAGE = (
    BASE + "/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/"
    "Paginas/imextractoejecucion.aspx"
)
_XLSX = re.compile(r'href="([^"]+\.xlsx)"', re.I)
_MONTHS = {
    m: i + 1
    for i, m in enumerate(
        [
            "enero",
            "febrero",
            "marzo",
            "abril",
            "mayo",
            "junio",
            "julio",
            "agosto",
            "septiembre",
            "octubre",
            "noviembre",
            "diciembre",
        ]
    )
}
_UNTIL = re.compile(r"HASTA EL (\d{1,2}) DE ([A-ZÁÉÍÓÚ]+)", re.I)


@dataclass
class Partida:
    nombre: str
    creditos: float  # creditos definitivos (o previsiones, en ingresos)
    obligaciones: float  # obligaciones reconocidas netas (o derechos reconocidos, en ingresos)
    pagos: float  # pagos realizados (o recaudacion neta, en ingresos)
    obligaciones_prev: float  # lo mismo del ano anterior, mismo mes

    @property
    def ejecucion(self) -> float | None:
        """Porcentaje ejecutado: obligaciones / creditos definitivos."""
        return round(100 * self.obligaciones / self.creditos, 1) if self.creditos > 0 else None

    @property
    def variacion(self) -> float | None:
        """Variacion interanual de las obligaciones (%)."""
        if self.obligaciones_prev <= 0:
            return None
        return round(100 * (self.obligaciones / self.obligaciones_prev - 1), 1)


@dataclass
class Extracto:
    url: str
    periodo: str  # AAAA-MM
    hasta: str  # "31 de agosto de 2026"
    gastos: list[Partida] = field(default_factory=list)
    inversiones: list[Partida] = field(default_factory=list)
    ingresos: list[Partida] = field(default_factory=list)
    total_gastos: Partida | None = None
    total_inversiones: Partida | None = None
    total_ingresos: Partida | None = None


def _num(v) -> float:
    if isinstance(v, int | float):
        return float(v)
    return 0.0  # la IGAE escribe "-" para cero


def _rows(ws) -> list[tuple]:
    return [r for r in ws.iter_rows(values_only=True) if r and isinstance(r[0], str) and r[0].strip()]


def _sheet(ws, income: bool = False) -> tuple[list[Partida], Partida | None]:
    items, total = [], None
    # Gastos: (1) creditos (2) comprometidos (3) obligaciones (4) pagos | (5)-(8) ano anterior.
    # Ingresos: (1) previsiones (2) derechos (3) recaudacion | (4)-(6) ano anterior.
    for row in _rows(ws)[1:]:
        name = " ".join(row[0].split())
        if len(row) < 8 or not any(isinstance(v, int | float) for v in row[1:5]):
            continue
        if income:
            p = Partida(name, _num(row[1]), _num(row[2]), _num(row[3]), _num(row[5]))
        else:
            p = Partida(name, _num(row[1]), _num(row[3]), _num(row[4]), _num(row[7]))
        if name.replace(" ", "").upper() == "TOTALES":
            total = p
        else:
            items.append(p)
    return items, total


def _period(title: str, year: str) -> tuple[str, str]:
    m = _UNTIL.search(title or "")
    if not m or not str(year).strip().isdigit():
        return "", ""
    month = _MONTHS.get(m.group(2).lower())
    if not month:
        return "", ""
    return f"{year}-{month:02d}", f"{int(m.group(1))} de {m.group(2).lower()} de {year}"


def parse(content: bytes, url: str = "") -> Extracto | None:
    """Extracto de un Excel de la IGAE ya descargado; None si no tiene la estructura esperada."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return None
    if "GTOS 001" not in wb.sheetnames:
        return None
    head = next(wb["GTOS 001"].iter_rows(min_row=2, max_row=2, values_only=True))
    periodo, hasta = _period(head[0], head[1])
    if not periodo:
        return None
    ex = Extracto(url=url, periodo=periodo, hasta=hasta)
    ex.gastos, ex.total_gastos = _sheet(wb["GTOS 001"])
    if "GTOS 002" in wb.sheetnames:
        ex.inversiones, ex.total_inversiones = _sheet(wb["GTOS 002"])
    if "ING 002" in wb.sheetnames:
        ex.ingresos, ex.total_ingresos = _sheet(wb["ING 002"], income=True)
    return ex


def latest_url(transport: httpx.BaseTransport | None = None) -> str:
    """URL del Excel mas reciente enlazado en la pagina del extracto (la IGAE lo pone el primero)."""
    with client(transport) as c:
        resp = c.get(PAGE)
    if resp.status_code != 200:
        return ""
    m = _XLSX.search(resp.text)
    return urljoin(BASE, m.group(1).replace("&amp;", "&")) if m else ""


def latest(transport: httpx.BaseTransport | None = None) -> Extracto | None:
    url = latest_url(transport)
    if not url:
        return None
    with client(transport) as c:
        resp = c.get(url)
    if resp.status_code != 200:
        return None
    return parse(resp.content, url)
